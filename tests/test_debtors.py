"""Fee debtors — who owes, summed live across terms, with contact and aging.

The bursar's chase-list. The property that matters: outstanding is computed
from invoices and payments at read time, never copied onto a new invoice, so
issued receipts are never altered by looking at the debtors list.
"""
import io

from openpyxl import load_workbook

from tests.conftest import headers


async def _bill(client, ah, ids, amount=50000, due="2026-09-01"):
    from app.models.academics import ClassArm
    async with ids["session_factory"]() as db:
        arm = await db.get(ClassArm, ids["arm_id"])
        class_id = arm.class_id
    cat = (await client.post("/api/fees/categories", headers=ah,
           json={"name": "School Fees"})).json()
    await client.post("/api/fees/structures", headers=ah, json={
        "class_id": class_id, "term_id": ids["term_id"],
        "category_id": cat["id"], "amount": amount})
    await client.post("/api/fees/invoices/generate", headers=ah, json={
        "arm_id": ids["arm_id"], "term_id": ids["term_id"], "due_date": due})
    return (await client.get("/api/fees/invoices", headers=ah,
            params={"term_id": ids["term_id"]})).json()


async def test_debtors_sum_live_and_rank_by_amount(ctx):
    client, ids = ctx
    ah = await headers(client, "admin@gss-ikeja.ng", "admin123")
    invs = await _bill(client, ah, ids, amount=50000)

    # one pays in full, one pays part, the rest owe the lot
    await client.post("/api/fees/payments", headers=ah, json={
        "invoice_id": invs[0]["id"], "amount": 50000, "method": "cash",
        "reference": "D-1", "paid_at": "2026-07-18"})
    await client.post("/api/fees/payments", headers=ah, json={
        "invoice_id": invs[1]["id"], "amount": 20000, "method": "cash",
        "reference": "D-2", "paid_at": "2026-07-18"})

    d = (await client.get("/api/fees/debtors", headers=ah)).json()
    n = len(ids["student_ids"])
    assert d["count"] == n - 1                       # the full payer is gone
    # three owe 50,000 and one owes 30,000
    assert d["total_outstanding"] == 50000 * (n - 2) + 30000
    # ranked: biggest debt first, the 30,000 part-payer last
    assert d["debtors"][0]["outstanding"] == 50000
    assert d["debtors"][-1]["outstanding"] == 30000
    # aging is populated from the due date
    assert all(row["oldest_due"] == "2026-09-01" for row in d["debtors"])
    assert all(row["days_overdue"] is not None for row in d["debtors"])


async def test_debtors_carry_guardian_contact(ctx):
    client, ids = ctx
    ah = await headers(client, "admin@gss-ikeja.ng", "admin123")
    invs = await _bill(client, ah, ids)

    # attach a guardian with a phone to the first student
    from app.core.security import hash_password
    from app.models.school import Role, User
    from app.models.student import Guardian
    async with ids["session_factory"]() as db:
        p = User(school_id=ids["school_id"], email="dad@x.ng",
                 hashed_password=hash_password("par12345"), role=Role.PARENT,
                 first_name="Baba", last_name="Musa", phone="08030000000")
        db.add(p)
        await db.flush()
        db.add(Guardian(school_id=ids["school_id"], parent_user_id=p.id,
                        student_id=ids["student_ids"][0]))
        await db.commit()

    d = (await client.get("/api/fees/debtors", headers=ah)).json()
    row = next(r for r in d["debtors"] if r["student_id"] == ids["student_ids"][0])
    assert row["guardian_name"] == "Baba Musa"
    assert row["guardian_phone"] == "08030000000"


async def test_debtors_filter_by_class(ctx):
    client, ids = ctx
    ah = await headers(client, "admin@gss-ikeja.ng", "admin123")
    await _bill(client, ah, ids)

    # the arm filter returns this arm's debtors; a different (empty) arm returns none
    same = (await client.get("/api/fees/debtors", headers=ah,
            params={"arm_id": ids["arm_id"]})).json()
    assert same["count"] == len(ids["student_ids"])
    none = (await client.get("/api/fees/debtors", headers=ah,
            params={"arm_id": "no-such-arm"})).json()
    assert none["count"] == 0


async def test_debtors_excel_export(ctx):
    client, ids = ctx
    ah = await headers(client, "admin@gss-ikeja.ng", "admin123")
    await _bill(client, ah, ids)

    r = await client.get("/api/fees/debtors.xlsx", headers=ah)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "debtors-" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header[0] == "Admission No" and header[-1] == "Outstanding (NGN)"
    # one data row per student, plus a spacer and a TOTAL row
    values = list(ws.values)
    assert any(row[-2] == "TOTAL" for row in values if row[-2])


async def test_debtors_is_finance_only(ctx):
    client, ids = ctx
    th = await headers(client, "teacher@gss-ikeja.ng", "teach123")
    assert (await client.get("/api/fees/debtors", headers=th)).status_code == 403
    assert (await client.get("/api/fees/debtors.xlsx", headers=th)).status_code == 403
